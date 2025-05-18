#!/usr/bin/python3
"""
[`pptx2audio.py`](https://github.com/deangelisdf/write2audiobook/blob/main/pptx2audio.py)

Convert a `.pptx` file to an M4B file.

Usage example:
    `python pptx2audio.py presentation.pptx en`
"""
import os
import logging
import pptx
import pptx.parts.image
from pptx import presentation, slide
from backend_audio import m4b
from frontend import input_tool

logging.basicConfig(level=logging.ERROR)
logger = logging.getLogger(__name__)

BACK_END_TTS = m4b.get_back_end_tts()
LANGUAGE = "it"

TOK_NUM_SLIDES = {"it":"Numero slide:", "en":"slide number: "}
TOK_SLIDE_NOTE = {"it":"Note:", "en":"Note:"}

def save_image_from_pptx(image:pptx.parts.image.Image, folder_path:str) -> None:
    """Save an image to a folder.

    Arguments:
        image: The image to save.
        folder_path: The folder to save the image to.
    """
    filename = f"{folder_path}/{image.filename}.{image.ext}"
    with open(filename, 'wb') as f:
        f.write(image.blob)

def __get_notes(note: slide.NotesSlide) -> str:
    print(note)
    return ""

def __extract_text_from_slide(slide_obj: slide.Slide,
                              language:str="it") -> str:
    text_slide, text_note = "", ""
    if slide_obj.has_notes_slide:
        text_note = __get_notes(slide_obj.notes_slide)
    for shape in slide_obj.shapes:
        if shape.has_text_frame:
            text_slide += shape.text
    text_slide = text_slide.strip()
    if len(text_note) != 0:
        text_slide += f"\n{TOK_SLIDE_NOTE[language]}\n{text_note}"
    return text_slide

def extract_pptx_text(path_pptx:str,
                      language:str="it") -> str:
    """Combine the text from each slide.

    Arguments:
        path_pptx (str): The PowerPoint presentation's file path.
        language: The language of the presentation's content.

    Returns:
        text_out: The presentation's text content.
    """
    text_out:str = ""
    p: presentation.Presentation = pptx.Presentation(path_pptx)
    for idx,s in enumerate(p.slides):
        text_out += f"\n{TOK_NUM_SLIDES[language]} {idx}\n"
        ts = __extract_text_from_slide(s)
        text_out += ts
        if len(ts) == 0:
            text_out += ""
    return text_out

def main():
    input_file_path, out_file_path, language = input_tool.get_sys_input(os.path.dirname(__file__))
    m4b.init(BACK_END_TTS)
    text=extract_pptx_text(input_file_path)
    m4b.generate_audio(text, out_file_path, lang=language, backend=BACK_END_TTS)

if __name__ == "__main__":
    main()
